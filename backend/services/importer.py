"""Deterministic + AI-assisted import logic (Phase C).

- Parse .csv / .xlsx (max 5 MB, max 20k rows).
- Propose source_column -> target_field mapping (LLM first, rule-based fallback).
- Commit: coerce types, validate with existing Pydantic Create schemas,
  resolve machine refs by code/name, bulk insert, skip duplicates/invalid.
- Telemetry imports NEVER auto-create incidents (historical rows).
"""
import csv
import io
import json
import re
from datetime import datetime
from typing import Any, Dict, List, Optional, Tuple
from uuid import UUID

from openpyxl import load_workbook
from pydantic import BaseModel, ValidationError

from backend.schemas.employee import EmployeeCreate
from backend.schemas.machine import MachineCreate
from backend.schemas.order import OrderCreate
from backend.schemas.telemetry import TelemetryCreate

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 20_000

TARGETS = ("telemetry", "machines", "employees", "orders")

# Whitelist: exactly the existing Create-schema fields plus the onboarding-added
# model columns (migration 008). Never accept factory_id from a file.
TARGET_FIELDS: Dict[str, List[str]] = {
    "telemetry": ["machine_code", "machine_id", "temperature", "vibration", "current", "rpm", "machine_status", "timestamp"],
    "machines": ["name", "machine_code", "machine_type", "department", "criticality", "location", "status", "health_status"],
    "employees": ["name", "employee_code", "role", "department", "experience_years", "skills", "certifications", "shift", "status", "availability"],
    "orders": ["order_number", "customer_name", "product", "quantity", "priority", "status", "deadline", "progress"],
}

# Model-only (onboarding) columns: present on the SQLAlchemy models but not on
# the Create schemas, so they are merged onto the model kwargs after validation.
EXTRA_FIELDS: Dict[str, Tuple[str, ...]] = {
    "machines": ("machine_code", "department", "criticality"),
    "employees": ("employee_code", "department", "experience_years"),
}

# Messy-header aliases for the deterministic fallback (Temp, Temp(C), Amp, Speed, Mach No, ...).
ALIASES: Dict[str, Dict[str, List[str]]] = {
    "telemetry": {
        "machine_code": ["mach no", "mach_no", "machine no", "machine code", "machinecode", "machine", "machine name", "machinename", "machine_id", "equipment", "asset"],
        "temperature": ["temp", "temp(c)", "temp_c", "temperature", "temperature_c", "heat", "degc"],
        "vibration": ["vib", "vibration", "vibe", "vib(mm/s)", "vibration_mm_s"],
        "current": ["amp", "amps", "current", "current_a", "amperage", "i(a)"],
        "rpm": ["speed", "rpm", "rev", "rotation", "rot speed"],
        "machine_status": ["status", "machine_status", "state", "run state", "run_state"],
        "timestamp": ["time", "timestamp", "date", "datetime", "recorded_at", "ts"],
    },
    "machines": {
        "name": ["name", "machine", "machine name", "machine_name", "equipment", "asset", "asset name"],
        "machine_code": ["machine id", "machine_id", "machine code", "machine_code", "machinecode", "asset id", "asset_id", "equipment id", "code"],
        "machine_type": ["type", "machine_type", "machine type", "category", "kind"],
        "department": ["department", "dept"],
        "criticality": ["criticality", "critical"],
        "location": ["location", "bay", "area", "site", "loc"],
        "status": ["status", "state"],
        "health_status": ["health", "health_status", "health status", "condition"],
    },
    "employees": {
        "name": ["name", "employee", "employee name", "employee_name", "full name", "full_name", "worker"],
        "employee_code": ["employee id", "employee_id", "emp id", "emp_id", "employee code", "employee_code", "employeecode", "emp code", "emp_code", "empcode", "id", "code"],
        "role": ["role", "job", "title", "position", "job title", "designation"],
        "department": ["department", "dept"],
        "experience_years": ["experience", "experience_years", "exp", "years", "years of experience"],
        "skills": ["skills", "skill", "competencies", "competency"],
        "certifications": ["certifications", "certs", "certificates", "cert", "certification"],
        "shift": ["shift", "shifts"],
        "status": ["status", "emp status"],
        "availability": ["availability", "available", "avail"],
    },
    "orders": {
        "order_number": ["order_number", "order no", "orderno", "order #", "order id", "order_id", "po", "order"],
        "customer_name": ["customer", "customer_name", "customer name", "client", "client name"],
        "product": ["product", "product name", "product_name", "item", "part", "description"],
        "quantity": ["qty", "quantity", "amount", "count", "units"],
        "priority": ["priority", "prio", "urgency"],
        "status": ["status", "order status", "state"],
        "deadline": ["deadline", "due", "due date", "due_date", "delivery", "delivery date", "delivery_date", "eta"],
        "progress": ["progress", "pct", "percent", "%", "completion"],
    },
}


class MappingProposal(BaseModel):
    mapping: Dict[str, Optional[str]]
    confidence: Dict[str, float]


def _norm(h: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", (h or "").strip().lower()).strip()


def _dedupe_headers(headers: List[str]) -> List[str]:
    seen: Dict[str, int] = {}
    out: List[str] = []
    for h in headers:
        base = (h or "").strip() or f"col_{len(out)}"
        if base not in seen:
            seen[base] = 1
            out.append(base)
        else:
            seen[base] += 1
            out.append(f"{base}_{seen[base]}")
    return out


def _csv_dialect(sample: str) -> str:
    first = (sample.splitlines() or [""])[0]
    if ";" in first and "," not in first:
        return ";"
    return ","


def parse_file(filename: str, content: bytes) -> Tuple[List[str], List[Dict[str, Any]]]:
    if len(content) > MAX_BYTES:
        raise ValueError(f"File too large ({len(content)} bytes). Max is 5 MB.")
    lower = (filename or "").lower()
    if lower.endswith(".csv"):
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise ValueError("CSV must be UTF-8 encoded.")
        reader = csv.DictReader(io.StringIO(text), delimiter=_csv_dialect(text))
        if not reader.fieldnames:
            raise ValueError("CSV has no header row.")
        headers = _dedupe_headers([h for h in reader.fieldnames if h is not None])
        rows = []
        for r in reader:
            d = {h: (v if v != "" else None) for h, v in r.items() if h is not None}
            # skip fully-empty rows (same as the xlsx path)
            if all(v is None for v in d.values()):
                continue
            rows.append(d)
            if len(rows) > MAX_ROWS:
                raise ValueError("File exceeds 20,000 rows.")
        return headers, rows
    if lower.endswith((".xlsx", ".xlsm")):
        # data_only=True: formulas are NEVER executed, only cached values read.
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        it = ws.iter_rows(values_only=True)
        try:
            header_row = next(it)
        except StopIteration:
            raise ValueError("Excel sheet is empty.")
        headers = _dedupe_headers(
            [str(h).strip() if h is not None and str(h).strip() != "" else f"col_{i}"
             for i, h in enumerate(header_row)]
        )
        rows = []
        for rec in it:
            d: Dict[str, Any] = {}
            for h, v in zip(headers, rec):
                d[h] = None if v == "" else v
            # skip fully-empty rows
            if all(v is None for v in d.values()):
                continue
            rows.append(d)
            if len(rows) > MAX_ROWS:
                raise ValueError("File exceeds 20,000 rows.")
        return headers, rows
    raise ValueError("Unsupported file type. Use .csv or .xlsx.")


def rule_based_mapping(target: str, headers: List[str]) -> Tuple[Dict[str, Optional[str]], Dict[str, float]]:
    aliases = ALIASES.get(target, {})
    norm_headers = {h: _norm(h) for h in headers}
    mapping: Dict[str, Optional[str]] = {}
    confidence: Dict[str, float] = {}
    used_targets = set()
    for h in headers:
        nh = norm_headers[h]
        best: Optional[str] = None
        best_conf = 0.0
        for field, keys in aliases.items():
            if field in used_targets:
                continue
            for k in keys:
                nk = _norm(k)
                if nh == nk:
                    best, best_conf = field, 1.0
                    break
                if nk and nk in nh:
                    if 0.7 > best_conf:
                        best, best_conf = field, 0.7
            if best_conf == 1.0:
                break
        # exact target-field name match always wins
        if nh in [f for f in TARGET_FIELDS.get(target, [])]:
            best, best_conf = nh, 1.0
        mapping[h] = best
        confidence[h] = best_conf
        if best:
            used_targets.add(best)
    return mapping, confidence


def llm_mapping(target: str, headers: List[str], samples: List[Dict[str, Any]]) -> Optional[Tuple[Dict[str, Optional[str]], Dict[str, float]]]:
    """Ask the LLM for a mapping; return None when unavailable/invalid."""
    try:
        from backend.services.llm import LLMUnavailable, complete_json  # lazy: keeps import light
    except Exception:
        return None
    try:
        result = complete_json(
            "You map spreadsheet columns to database fields. Output JSON {mapping: {source_column: target_field|null}, confidence: {source_column: 0..1}}. "
            f"Allowed target fields for '{target}': {TARGET_FIELDS[target]}. Use null when no field fits. JSON ONLY.",
            {"target": target, "headers": headers, "sample_rows": samples[:5]},
            MappingProposal,
        )
        mapping = {h: result.mapping.get(h) for h in headers}
        # Validate: drop unknown fields.
        allowed = set(TARGET_FIELDS[target])
        for h, f in list(mapping.items()):
            if f is not None and f not in allowed:
                mapping[h] = None
        conf = {h: float(result.confidence.get(h, 0.5)) for h in headers}
        return mapping, conf
    except Exception:
        return None


def propose_mapping(target: str, headers: List[str], samples: List[Dict[str, Any]]) -> Tuple[Dict[str, Optional[str]], Dict[str, float], str]:
    llm = llm_mapping(target, headers, samples)
    if llm is not None:
        return llm[0], llm[1], "ai"
    m, c = rule_based_mapping(target, headers)
    return m, c, "rule-based"


def _coerce_number(v: Any) -> Optional[float]:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "")
    try:
        return float(s)
    except ValueError:
        return None


def _coerce_int(v: Any) -> Optional[int]:
    n = _coerce_number(v)
    return None if n is None else int(n)


def _coerce_datetime(v: Any) -> Optional[datetime]:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v
    s = str(v).strip()
    # Excel serial dates sometimes arrive as numbers via openpyxl data_only=False; read_only gives datetime already.
    try:
        # fromisoformat handles "2026-05-01", "2026-05-01T10:00:00", with space separator too.
        return datetime.fromisoformat(s.replace("Z", "+00:00"))
    except ValueError:
        pass
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y", "%d.%m.%Y"):
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            continue
    return None


def describe_mapping(target: str, headers: List[str],
                     mapping: Dict[str, Optional[str]],
                     confidence: Dict[str, float],
                     source: str) -> Tuple[List[Dict[str, Any]], List[str]]:
    """Per-column proposal details + ambiguity warnings for the review UI.

    Returns (details, warnings) where each detail is
    {source_column, destination_field, confidence, reason}.
    """
    aliases = ALIASES.get(target, {})
    norm_headers = {h: _norm(h) for h in headers}
    details: List[Dict[str, Any]] = []
    warnings: List[str] = []
    claimed_by: Dict[str, str] = {}
    for h in headers:
        fld = mapping.get(h)
        conf = float(confidence.get(h, 0.0))
        if fld is None:
            # Was there a field this header wanted but an earlier column took?
            nh = norm_headers[h]
            want: Optional[str] = None
            for field, keys in aliases.items():
                if any(_norm(k) == nh or (_norm(k) and _norm(k) in nh) for k in keys):
                    want = field
                    break
            if want is not None and want in claimed_by:
                reason = f"also looks like '{want}' (already mapped from '{claimed_by[want]}')"
                warnings.append(f"Ambiguous: '{h}' {reason}; left unmapped.")
            else:
                reason = "no match — map manually or leave ignored"
            details.append({"source_column": h, "destination_field": None,
                            "confidence": 0.0, "reason": reason})
            continue
        claimed_by.setdefault(fld, h)
        if source == "ai":
            reason = "AI proposal"
        elif conf >= 1.0:
            reason = "exact header match"
        else:
            reason = "partial match — please verify"
            warnings.append(f"Verify mapping: '{h}' → '{fld}' (confidence {conf:.2f}).")
        details.append({"source_column": h, "destination_field": fld,
                        "confidence": conf, "reason": reason})
    unmapped = [h for h in headers if not mapping.get(h)]
    if unmapped:
        warnings.append(f"{len(unmapped)} column(s) unmapped: {', '.join(unmapped[:8])}"
                        + ("…" if len(unmapped) > 8 else "") + ". They will be ignored.")
    return details, warnings


def extract_extras(target: str, mapped: Dict[str, Any]) -> Tuple[Dict[str, Any], List[str]]:
    """Onboarding-added model columns (no Create-schema equivalent).

    Returns (extras, warnings). All values are treated as plain data.
    """
    extras: Dict[str, Any] = {}
    warnings: List[str] = []
    if target == "employees":
        code = mapped.get("employee_code")
        extras["employee_code"] = str(code).strip() or None if code not in (None, "") else None
        dept = mapped.get("department")
        extras["department"] = str(dept).strip() or None if dept not in (None, "") else None
        raw_exp = mapped.get("experience_years")
        if raw_exp in (None, ""):
            extras["experience_years"] = None
        else:
            n = _coerce_int(raw_exp)
            if n is None:
                warnings.append(f"ignoring bad experience_years '{raw_exp}'")
                extras["experience_years"] = None
            else:
                extras["experience_years"] = n
    elif target == "machines":
        code = mapped.get("machine_code")
        extras["machine_code"] = str(code).strip().upper() or None if code not in (None, "") else None
        dept = mapped.get("department")
        extras["department"] = str(dept).strip() or None if dept not in (None, "") else None
        crit = mapped.get("criticality")
        extras["criticality"] = str(crit).strip().upper() or None if crit not in (None, "") else None
    return extras, warnings


def build_import_row(target: str, mapped: Dict[str, Any]) -> Tuple[Optional[Dict[str, Any]], Optional[str], List[str]]:
    """Coerce + validate one mapped row. Returns (model_kwargs, error, warnings).

    model_kwargs are ready for Model(**kwargs); error is a human message.
    Rejects empty natural keys and uncoercible numerics that the schemas
    would otherwise silently default.
    """
    if target == "orders":
        if not str(mapped.get("order_number") or "").strip():
            return None, "missing order_number", []
        raw_qty = mapped.get("quantity")
        if raw_qty not in (None, "") and _coerce_int(raw_qty) is None:
            return None, f"bad quantity: '{raw_qty}'", []
        raw_prog = mapped.get("progress")
        if raw_prog not in (None, "") and _coerce_number(raw_prog) is None:
            return None, f"bad progress: '{raw_prog}'", []
    elif target == "machines":
        if not str(mapped.get("name") or "").strip():
            return None, "missing machine name", []
    elif target == "employees":
        if not str(mapped.get("name") or "").strip():
            return None, "missing employee name", []
        if not str(mapped.get("role") or "").strip():
            return None, f"missing role for '{str(mapped.get('name') or '').strip()}'", []
    obj, err = build_validated_row(target, mapped)
    if err or obj is None:
        return None, err or "validation failed", []
    kwargs: Dict[str, Any] = dict(obj.model_dump())
    warnings: List[str] = []
    if target in EXTRA_FIELDS:
        extras, w = extract_extras(target, mapped)
        kwargs.update(extras)
        warnings.extend(w)
    return kwargs, None, warnings


def _machine_code_of(name: Optional[str], explicit: Optional[Any]) -> Optional[str]:
    if explicit not in (None, ""):
        return str(explicit).strip().upper() or None
    m = re.match(r"^(M-\d{3})", (name or "").strip(), re.IGNORECASE)
    return m.group(1).upper() if m else None


def load_import_context(db, target: str) -> Dict[str, Any]:
    """Scoped read-only lookups for duplicate detection + machine resolution.

    All queries go through the caller's scoped session, so another company's
    rows are invisible here exactly as in every other ORM read.
    """
    from backend.models.models import Employee, Machine, Order
    ctx: Dict[str, Any] = {"machines": [], "code_to_id": {}, "name_to_id": {},
                           "orders": set(), "machine_names": set(), "machine_codes": set(),
                           "employees": set(), "employee_codes": set()}
    machines = db.query(Machine).all()
    ctx["machines"] = machines
    for m in machines:
        ctx["name_to_id"][(m.name or "").strip().lower()] = m.id
        code = (getattr(m, "machine_code", None) or "").strip().upper()
        if code:
            ctx["code_to_id"][code] = m.id
        pm = re.match(r"^(M-\d{3})", (m.name or "").strip(), re.IGNORECASE)
        if pm:
            ctx["code_to_id"].setdefault(pm.group(1).upper(), m.id)
    if target == "orders":
        ctx["orders"] = {r[0] for r in db.query(Order.order_number).all()}
    elif target == "machines":
        ctx["machine_names"] = {m.name for m in machines}
        ctx["machine_codes"] = set(ctx["code_to_id"])
    elif target == "employees":
        ctx["employees"] = {(e.name, e.role) for e in db.query(Employee.name, Employee.role).all()}
        ctx["employee_codes"] = {(e.employee_code or "").strip() for e in
                                 db.query(Employee.employee_code).all()
                                 if (e.employee_code or "").strip()}
    return ctx


def resolve_import_machine_id(mapped: Dict[str, Any], ctx: Dict[str, Any]) -> Tuple[Any, Optional[str]]:
    """Resolve a telemetry row's machine. Returns (uuid|None, error|None)."""
    raw = mapped.get("machine_code", mapped.get("machine_id"))
    if raw in (None, ""):
        return None, "missing machine reference (machine_code)"
    s = raw.strip() if isinstance(raw, str) else str(raw)
    su = s.upper()
    if su in ctx["code_to_id"]:
        return ctx["code_to_id"][su], None
    m = re.match(r"^(M-\d{3})", s.strip(), re.IGNORECASE)
    if m and m.group(1).upper() in ctx["code_to_id"]:
        return ctx["code_to_id"][m.group(1).upper()], None
    if s.strip().lower() in ctx["name_to_id"]:
        return ctx["name_to_id"][s.strip().lower()], None
    try:
        uid = UUID(s)
    except (ValueError, AttributeError):
        return None, f"unknown machine '{s}'"
    # A UUID string still has to belong to this factory's machines.
    if uid in set(ctx["name_to_id"].values()) or uid in set(ctx["code_to_id"].values()):
        return uid, None
    known = {m.id for m in ctx["machines"]}
    if uid in known:
        return uid, None
    return None, f"unknown machine '{s}'"


def prepare_import_rows(target: str, rows: List[Dict[str, Any]],
                        confirmed: Dict[str, Optional[str]],
                        ctx: Dict[str, Any]) -> Tuple[List[Tuple[int, Dict[str, Any]]],
                                                     List[Dict[str, Any]],
                                                     List[Dict[str, Any]], List[str]]:
    """Validate every row in memory. No database writes.

    Returns (valid, invalid, duplicates, warnings):
    - valid: [(excel_row_number, model_kwargs)]
    - invalid: [{row, error}]
    - duplicates: [{row, key}]
    """
    valid: List[Tuple[int, Dict[str, Any]]] = []
    invalid: List[Dict[str, Any]] = []
    duplicates: List[Dict[str, Any]] = []
    warnings: List[str] = []
    seen_orders = set()
    seen_machine_names = set()
    seen_machine_codes = set()
    seen_employees = set()
    seen_employee_codes = set()

    for offset, raw_row in enumerate(rows):
        idx = offset + 2  # 1-based + header row
        mapped = {fld: raw_row.get(src) for src, fld in confirmed.items() if fld}
        for w in extract_extras(target, mapped)[1]:
            warnings.append(f"Row {idx}: {w}")
        if target == "telemetry" and mapped.get("machine_id") in (None, ""):
            mid, err = resolve_import_machine_id(mapped, ctx)
            if err:
                invalid.append({"row": idx, "error": err})
                continue
            mapped["machine_id"] = mid
        elif target == "telemetry":
            try:
                UUID(str(mapped["machine_id"]))
            except (ValueError, AttributeError):
                mid, err = resolve_import_machine_id({"machine_code": mapped["machine_id"]}, ctx)
                if err:
                    invalid.append({"row": idx, "error": err})
                    continue
                mapped["machine_id"] = mid
            else:
                uid = UUID(str(mapped["machine_id"]))
                known = {m.id for m in ctx["machines"]}
                if uid not in known:
                    invalid.append({"row": idx, "error": f"unknown machine '{mapped['machine_id']}'"})
                    continue
        kwargs, err, _w = build_import_row(target, mapped)
        if err or kwargs is None:
            invalid.append({"row": idx, "error": err or "validation failed"})
            continue
        if target == "orders":
            key = kwargs["order_number"]
            if key in ctx["orders"] or key in seen_orders:
                duplicates.append({"row": idx, "key": key})
                continue
            seen_orders.add(key)
        elif target == "machines":
            code = _machine_code_of(kwargs.get("name"), kwargs.get("machine_code"))
            if kwargs.get("name") in ctx["machine_names"] or kwargs.get("name") in seen_machine_names \
                    or (code and (code in ctx["machine_codes"] or code in seen_machine_codes)):
                duplicates.append({"row": idx, "key": kwargs.get("name")})
                continue
            seen_machine_names.add(kwargs.get("name"))
            if code:
                seen_machine_codes.add(code)
        elif target == "employees":
            ecode = (kwargs.get("employee_code") or "").strip() if kwargs.get("employee_code") else ""
            ekey = (kwargs.get("name"), kwargs.get("role"))
            if (ecode and (ecode in ctx["employee_codes"] or ecode in seen_employee_codes)) \
                    or (ekey in ctx["employees"] or ekey in seen_employees):
                duplicates.append({"row": idx, "key": ecode or f"{ekey[0]} / {ekey[1]}"})
                continue
            seen_employees.add(ekey)
            if ecode:
                seen_employee_codes.add(ecode)
        # telemetry: no dedup — historical rows insert as-is (existing behavior)
        valid.append((idx, kwargs))
    return valid, invalid, duplicates, warnings


def _jsonable(value: Any) -> Any:
    from datetime import date
    if isinstance(value, UUID):
        return str(value)
    if isinstance(value, (datetime, date)):
        try:
            return value.isoformat()
        except Exception:
            return str(value)
    if isinstance(value, dict):
        return {k: _jsonable(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [_jsonable(v) for v in value]
    return value


def _require_import_scope(db) -> None:
    from backend.db.scoping import SCOPE_KEY
    try:
        scope = db.info.get(SCOPE_KEY)
    except Exception:
        scope = None
    if not scope:
        from fastapi import HTTPException
        raise HTTPException(status_code=400, detail="No active factory. Send the X-Factory-Id header.")


def run_import_preview(db, target: str, rows: List[Dict[str, Any]],
                       confirmed: Dict[str, Optional[str]],
                       max_examples: int = 25) -> Dict[str, Any]:
    """Validate + classify every row. Reads only — never writes."""
    _require_import_scope(db)
    ctx = load_import_context(db, target)
    valid, invalid, duplicates, warnings = prepare_import_rows(target, rows, confirmed, ctx)
    return {
        "target": target,
        "row_count": len(rows),
        "valid_count": len(valid),
        "invalid_count": len(invalid),
        "duplicate_count": len(duplicates),
        "valid_rows": [_jsonable(kw) for _, kw in valid[:max_examples]],
        "invalid_rows": invalid[:20],
        "duplicate_rows": duplicates[:20],
        "warnings": warnings[:20],
    }


def run_import_confirm(db, target: str, rows: List[Dict[str, Any]],
                       confirmed: Dict[str, Optional[str]]) -> Dict[str, Any]:
    """Validate everything, then insert in ONE transaction. Nothing on failure.

    Row-level invalid/duplicates are skipped + reported; only a commit-level
    database error rolls back the whole batch (inserted = 0).
    """
    _require_import_scope(db)
    from backend.models.models import Employee, Machine, MachineTelemetry, Order
    ctx = load_import_context(db, target)
    valid, invalid, duplicates, warnings = prepare_import_rows(target, rows, confirmed, ctx)
    errors: List[Dict[str, Any]] = list(invalid)
    if not valid:
        return {
            "target": target,
            "inserted": 0, "imported_count": 0,
            "skipped_duplicates": len(duplicates), "skipped_count": len(duplicates),
            "failed": len(invalid), "errors": errors[:20], "warnings": warnings[:20],
            "note": "Telemetry imports never auto-create incidents (historical rows)."
                    if target == "telemetry" else None,
        }
    models = {"orders": Order, "machines": Machine,
              "employees": Employee, "telemetry": MachineTelemetry}
    model = models[target]
    try:
        for _, kwargs in valid:
            db.add(model(**kwargs))
        db.commit()
    except Exception as e:
        db.rollback()
        errors = errors + [{"row": "commit", "error": str(e)}]
        return {
            "target": target,
            "inserted": 0, "imported_count": 0,
            "skipped_duplicates": len(duplicates), "skipped_count": len(duplicates),
            "failed": len(valid) + len(invalid), "errors": errors[:20], "warnings": warnings[:20],
            "note": "Commit failed — nothing was written.",
        }
    return {
        "target": target,
        "inserted": len(valid), "imported_count": len(valid),
        "skipped_duplicates": len(duplicates), "skipped_count": len(duplicates),
        "failed": len(invalid), "errors": errors[:20], "warnings": warnings[:20],
        "note": "Telemetry imports never auto-create incidents (historical rows)."
                if target == "telemetry" else None,
    }


def _coerce_json_dict(v: Any) -> Optional[Dict[str, Any]]:
    if v is None or v == "":
        return None
    if isinstance(v, dict):
        return v
    if isinstance(v, list):
        return {"items": v}
    s = str(v).strip()
    try:
        parsed = json.loads(s)
        if isinstance(parsed, dict):
            return parsed
        if isinstance(parsed, list):
            return {"items": parsed}
        return {"value": parsed}
    except (json.JSONDecodeError, ValueError):
        pass
    # comma-separated skills string -> {"items": [...]}
    parts = [p.strip() for p in re.split(r"[;,|]", s) if p.strip()]
    if len(parts) > 1:
        return {"items": parts}
    return {"value": s} if s else None


def build_validated_row(target: str, mapped: Dict[str, Any]) -> Tuple[Optional[BaseModel], Optional[str]]:
    """Coerce + validate one mapped row. Returns (obj, error)."""
    try:
        if target == "telemetry":
            data: Dict[str, Any] = {}
            if mapped.get("machine_id"):
                try:
                    data["machine_id"] = UUID(str(mapped["machine_id"]))
                except ValueError:
                    return None, f"bad machine_id '{mapped['machine_id']}'"
            for f in ("temperature", "vibration", "current", "rpm"):
                if mapped.get(f) is not None:
                    n = _coerce_number(mapped[f])
                    if mapped[f] not in (None, "") and n is None:
                        return None, f"bad number for {f}: '{mapped[f]}'"
                    data[f] = n
            if mapped.get("machine_status") is not None:
                data["machine_status"] = str(mapped["machine_status"]).strip().upper() or None
            if mapped.get("timestamp") is not None:
                dt = _coerce_datetime(mapped["timestamp"])
                if mapped["timestamp"] not in (None, "") and dt is None:
                    return None, f"bad timestamp: '{mapped['timestamp']}'"
                data["timestamp"] = dt
            obj = TelemetryCreate(**data)
            return obj, None
        if target == "machines":
            obj = MachineCreate(
                name=str(mapped.get("name") or "").strip(),
                machine_type=(str(mapped.get("machine_type")).strip() if mapped.get("machine_type") else None),
                location=(str(mapped.get("location")).strip() if mapped.get("location") else None),
                status=(str(mapped.get("status")).strip().upper() if mapped.get("status") else "OPERATIONAL"),
                health_status=(str(mapped.get("health_status")).strip().upper() if mapped.get("health_status") else "GOOD"),
            )
            return obj, None
        if target == "employees":
            obj = EmployeeCreate(
                name=str(mapped.get("name") or "").strip(),
                role=str(mapped.get("role") or "").strip(),
                skills=_coerce_json_dict(mapped.get("skills")),
                certifications=_coerce_json_dict(mapped.get("certifications")),
                shift=(str(mapped.get("shift")).strip() if mapped.get("shift") else None),
                status=(str(mapped.get("status")).strip().upper() if mapped.get("status") else "ACTIVE"),
                availability=(str(mapped.get("availability")).strip().upper() if mapped.get("availability") else "AVAILABLE"),
            )
            return obj, None
        if target == "orders":
            qty = _coerce_int(mapped.get("quantity")) if mapped.get("quantity") not in (None, "") else 0
            prog = _coerce_number(mapped.get("progress")) if mapped.get("progress") not in (None, "") else 0.0
            dl = _coerce_datetime(mapped.get("deadline")) if mapped.get("deadline") not in (None, "") else None
            if mapped.get("deadline") not in (None, "") and dl is None:
                return None, f"bad deadline: '{mapped['deadline']}'"
            obj = OrderCreate(
                order_number=str(mapped.get("order_number") or "").strip(),
                customer_name=(str(mapped.get("customer_name")).strip() if mapped.get("customer_name") else None),
                product=(str(mapped.get("product")).strip() if mapped.get("product") else None),
                quantity=qty or 0,
                priority=(str(mapped.get("priority")).strip().upper() if mapped.get("priority") else "NORMAL"),
                status=(str(mapped.get("status")).strip().upper() if mapped.get("status") else "PENDING"),
                deadline=dl,
                progress=prog or 0.0,
            )
            return obj, None
    except ValidationError as e:
        return None, str(e.errors()[0].get("msg", e)) if e.errors() else str(e)
    except Exception as e:
        return None, str(e)
    return None, f"unknown target '{target}'"
